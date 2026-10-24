#!/usr/bin/env python3
"""Write data/angelone-seed.xlsx: the AngelOne-1 tab for the Brand Feeds sheet.

Every number in these posts was read off angelone.in, its llms.txt, its Tag Manager
container, the @angelone Instagram grid and the Angel One YouTube channel on
25 September 2026. Import the file as a tab named AngelOne-1 in the Brand Feeds
sheet, or sync from it directly:  python3 tools/sync_sheet.py --file data/angelone-seed.xlsx
"""
import os, sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheet_schema import BRAND_KEYS, COLUMNS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BRAND = {
    'name': 'Angel One', 'host': 'angelone.in', 'products': 16149, 'unit': 'instrument pages',
    'logo': 'assets/ao-logo-tile.png', 'colors': '#2E51FF, #2A394E, #FF7300, #F1F4FD',
    'type_sample': 'AngelOne', 'type_name': 'Barlow headings, Roboto body',
    'type_face': "'Barlow', 'Roboto', system-ui, sans-serif",
    'note': 'v1: angelone.in, @angelone Instagram and the Angel One YouTube channel, 25 Sep 2026',
}
MG = ('meta', '#0866FF')

def P(pid, loc, typ, agent, **kw):
    d = {'post_id': pid, 'status': 'live', 'location': loc, 'type': typ, 'agent': agent}
    for k, v in kw.items():
        if k == 'images': v = '\n'.join(v)
        elif k == 'rows': k, v = 'table_rows', '\n'.join(f'{a} | {b}' for a, b in v)
        elif k == 'tags': v = ', '.join(v)
        elif k == 'extra': k, v = 'extra_json', json.dumps(v, ensure_ascii=False)
        elif k == 'cta':
            d['cta_label'] = v[0]
            if len(v) > 1: d['cta_icon'], d['cta_color'] = v[1], v[2]
            continue
        d[k] = v
    return d

FD = 'feed, deck'
POSTS = [
 P('ao-01', FD, 'image', 'creative', time='9 min', campaign='IPO Bulletin', image_ratio='4:5',
   title="Run this week's IPO Bulletins as one carousel",
   description='A-One Steels, ArMee Infotech and Varmora Granito share one layout: revenue and PAT as 3D stacks, FY24 against FY26. Approving publishes all three to Instagram as one carousel.',
   images=['assets/ao-ig-bulletin-aone.jpg', 'assets/ao-ig-bulletin-armee.jpg', 'assets/ao-ig-bulletin-varmora.jpg'],
   cta=('Publish Campaign',) + MG),
 P('ao-02', FD, 'table', 'performance', time='14 min', badge='Tracking', badge_tone='warn',
   title='Give each account opening one conversion',
   description='Tag Manager on angelone.in starts 7 Meta pixels and 7 Google Ads conversion IDs, so one sign-up can be counted many times over. Approving maps each goal to one pixel and one conversion.',
   rows=[('Meta pixels initialised', '7'), ('Google Ads conversion IDs', '7'),
         ('Also tagged', 'Microsoft Ads, Floodlight, Quora, LinkedIn'), ('Container', 'GTM-NNZ26S')],
   cta=('Audit conversions',)),
 P('ao-03', FD, 'table', 'storefront', deck_column='catalog', time='22 min', badge='Renamed', badge_tone='warn',
   title='Retitle the Tata Motors page for the demerger',
   description="Your homepage links 'Tata Motors share price' to a page that now shows Tata Motors Passenger Vehicles. Approving retitles it and links the commercial vehicles page beside it.",
   rows=[('Homepage link', 'Tata Motors share price'), ('Page title', 'Tata Motors Share Price Today'),
         ('What the page shows', 'Tata Motors Passenger Vehicles (TMPV)'),
         ('Commercial vehicles', 'its own page, not linked')],
   cta=('Apply Update',)),
 P('ao-04', FD, 'image', 'creative', time='31 min', campaign='IPO Shorts', image_ratio='square',
   title="Make '5 things to know' the IPO Shorts format",
   description='Your NSE IPO short, five things to know before applying, has 51,000 views. The IPO shorts around it sit near 10,000. Approving recuts the next three IPO explainers into the same five-point frame.',
   images=['assets/ao-vid-nse-ipo-poster.jpg'],
   extra={'video': 'assets/ao-vid-nse-ipo.mp4', 'poster': 'assets/ao-vid-nse-ipo-poster.jpg'},
   cta=('Approve format',)),
 P('ao-05', FD, 'table', 'storefront', time='38 min', badge='Mismatch', badge_tone='bad',
   title='Pick one line for delivery brokerage',
   description='The pricing page banner says ₹0 brokerage on stock investments. The table under it charges up to ₹20 an order after 30 days. Approving aligns every page to the rate you confirm.',
   rows=[('Pricing page banner', '₹0 on stock investments'),
         ('Pricing table, stock investments', 'lower of ₹20 or 0.1% after 30 days'),
         ('Open Demat page', 'Rs 0 on Equity Delivery')],
   cta=('Align pricing',)),
 P('ao-06', FD, 'image', 'performance', time='44 min', campaign='SuperApp for Smart Investing', image_ratio='square',
   title="Refresh the SuperApp film's end card",
   description='The film closes on 2 Crore+ Indians, sourced to January 2024, and 25 years. The site now says 3.8 Cr+ and 30 years, and the logo has changed. Approving updates the card and relaunches it.',
   images=['assets/ao-superapp-endcard.jpg'],
   cta=('Update and relaunch',)),
 P('ao-07', FD, 'image', 'creative', time='1 hr', campaign='Case studies', image_ratio='square',
   title='Turn the Hindi case study into a series',
   description='The Ather story, the only Hindi title in your last 30 videos, has 1.6 lakh views, 2.5 times the next best. Approving schedules Hindi cuts of the Adani airports and Airtel vs Jio case studies.',
   images=['assets/ao-yt-ather.jpg', 'assets/ao-yt-adani.jpg', 'assets/ao-yt-jio-airtel.jpg'],
   cta=('Schedule series',)),
 P('ao-08', 'feed', 'connect', 'orchestrator', connect='wordpress'),
 P('ao-09', FD, 'table', 'visibility', time='1 hr', badge='Out of date', badge_tone='bad',
   title='Correct the brokerage line AI models read',
   description='Your llms.txt tells AI models stock trading has zero brokerage on delivery. Your pricing table charges up to ₹20 an order after 30 days. Approving rewrites the line to match.',
   rows=[('llms.txt says', 'zero brokerage on delivery'),
         ('Pricing table says', 'lower of ₹20 or 0.1%, minimum ₹5'),
         ('llms.txt last updated', 'May 2026')],
   cta=('Rewrite llms.txt',)),
 P('ao-10', FD, 'image', 'creative', time='2 hr', campaign='Podcast', image_ratio='square',
   title='Hold the ₹1 crore cover for a compliance read',
   description='This cover says anyone watching the podcast can make ₹1 crore, which reads as a promised return. Approving sends it to compliance and runs the question-led cover in its slot.',
   images=['assets/ao-ig-podcast-crore.jpg', 'assets/ao-ig-podcast-winning.jpg'],
   cta=('Send to compliance',)),
 P('ao-11', FD, 'image', 'storefront', deck_column='catalog', time='2 hr', image_ratio='square',
   title='Tag each app render to the product it sells',
   description='These four renders show zero-commission funds, portfolio advisory, fast orders and TradingView. Approving tags each to its product page, so every listing opens on the screen it is selling.',
   images=['assets/ao-app-mf.jpg', 'assets/ao-app-advisory.jpg', 'assets/ao-app-orders.jpg', 'assets/ao-app-tradingview.jpg'],
   cta=('Tag and publish',)),
 P('ao-12', FD, 'image', 'performance', time='3 hr', campaign='Monthly savings made easy', image_ratio='square',
   title='Put the ₹100 start in the SIP ad',
   description='The SIP banner says Start SIP with no amount. The app and the start-sip page both lead with ₹100. Approving adds the ₹100 start to the Meta and Google versions before they spend.',
   images=['assets/ao-og-sip.jpg'],
   cta=('Launch campaign',) + MG),
 P('ao-13', FD, 'chart', 'visibility', time='3 hr',
   title='Put Angel One first in the stock-page descriptions',
   description="Every stock page description ends with 'at Angel One'. In a 25-page sample the median runs 179 characters and search shows about 160, so your name is what gets cut. Approving moves it up front.",
   extra={'chart': {'title': 'Meta description length, 25 stock pages sampled', 'max': 220, 'unit': '',
                    'note': 'Every page sampled runs past this',
                    'rows': [{'label': 'Longest', 'value': 207}, {'label': 'Median', 'value': 179, 'brand': True},
                             {'label': 'Shortest', 'value': 161}, {'label': 'Shown in search', 'value': 160}]}},
   cta=('Rewrite descriptions',)),
 P('ao-14', FD, 'table', 'storefront', time='4 hr', badge='Mismatch', badge_tone='bad',
   title='Use one customer count across the site',
   description='The site quotes four customer counts: 3.8 Cr+ on the homepage, 3.5 Cr+ on pricing, 2.5 Cr+ on mutual funds and 1.5 crore on the app page. Approving updates every page to 3.8 Cr+.',
   rows=[('Homepage and About', '3.8 Cr+'), ('Pricing page', '3.5 Cr+'),
         ('Mutual funds page', '2.5 Cr+'), ('App page intro', '1.5 crore')],
   cta=('Apply Update',)),
 P('ao-15', FD, 'text', 'creative', time='4 hr',
   title='Shoot real India, phone in hand',
   description="The SuperApp film is set in a lift, a bazaar and a village lane, one feature per scene. Every short is the same warm studio. Approving makes 'real place, phone in hand' the brief for the next shoot.",
   cta=('Approve brief',)),
 P('ao-16', FD, 'table', 'storefront', deck_column='catalog', time='5 hr', badge='Stale', badge_tone='warn',
   title='Move stale pages out of Upcoming IPO',
   description='The footer lists six IPOs as upcoming. boAt still cites a January 2022 filing and Snapdeal already shows a price band. Approving refreshes the list from issues actually opening.',
   rows=[('boAt', 'DRHP of 26 Jan 2022, awaiting approval'), ('Oyo', 'funding story ends in 2021'),
         ('Snapdeal', 'price band already set, ₹30 to ₹32'),
         ('Reliance Jio, PhonePe, Zepto', 'every date To be announced')],
   cta=('Refresh the list',)),
 P('ao-17', FD, 'image', 'creative', time='5 hr', campaign='SuperApp for Smart Investing', image_ratio='square',
   title='Cut the SuperApp film to one feature per Reel',
   description='The 45-second film packs low-network trading, buying from the chart and PIN login into one spot. Approving cuts a 15-second Reel for each and publishes them to Instagram and YouTube.',
   images=['assets/ao-vid-superapp-poster.jpg'],
   extra={'video': 'assets/ao-vid-superapp.mp4', 'poster': 'assets/ao-vid-superapp-poster.jpg'},
   cta=('Publish Reel',) + MG),
 P('ao-18', FD, 'image', 'performance', time='6 hr', campaign='3.8 Crore customers', image_ratio='square',
   title='Run the 3.8 Crore banner as the proof ad',
   description='The 3.8 Crore customers banner is already made and live on your About page. Approving runs it on Meta for people who opened the demat page but never finished an account.',
   images=['assets/ao-3-8-crore.jpg'],
   cta=('Launch campaign',) + MG),
 P('ao-19', FD, 'image', 'storefront', time='7 hr', image_ratio='square',
   title='Replace the share image on angelone.in',
   description='When angelone.in is shared on WhatsApp or X, the preview is a SIP banner uploaded in August 2024. Approving swaps in a frame from the current SuperApp film as the homepage share image.',
   images=['assets/ao-credits-hero.jpg'],
   cta=('Apply Update',)),
 P('ao-20', FD, 'image', 'creative', time='7 hr', campaign='Apply or Skip', image_ratio='square',
   title='Make Apply or Skip the cover for every mainboard IPO',
   description="Apply or Skip puts the viewer's decision in the headline, and the Elevate Campuses cover is the clearest frame in your grid. Approving makes it the cover template for the next three mainboard IPOs.",
   images=['assets/ao-ig-apply-skip.jpg'],
   cta=('Approve template',)),
 P('ao-21', FD, 'table', 'visibility', time='8 hr', badge='Needs you', badge_tone='info',
   title='Point llms.txt at the live IPO and fund pages',
   description='llms.txt points AI models at 98 pages. None is one of your 1,292 IPO pages or 7,266 fund pages. Approving adds the IPOs open this week and the most-viewed funds, refreshed every Monday.',
   rows=[('Pages in llms.txt', '98'), ('IPO pages on the site', '1,292'),
         ('Fund pages on the site', '7,266'), ('Of those in llms.txt', '0')],
   cta=('Update llms.txt',)),
 P('ao-22', FD, 'image', 'creative', time='9 hr', campaign='Angel One Markets', image_ratio='square',
   title='Run market news as one 3D scene',
   description='The Policybazaar post turns a 34% fall into a single frame: stacked papers, a red ₹30,000 Cr slip, the city behind glass. Approving makes it the template for daily market-news posts.',
   images=['assets/ao-ig-news-policybazaar.jpg'],
   cta=('Approve template',)),
 P('ao-23', FD, 'image', 'performance', time='10 hr', campaign='Invest in the next big IPO', image_ratio='square',
   title='Retarget IPO page visitors while the issue is open',
   description='Your 1,292 IPO pages are where IPO intent lands. Approving builds a Meta and Google audience from their visitors and shows them this frame until bidding closes.',
   images=['assets/ao-app-ipo.jpg'],
   cta=('Launch campaign',) + MG),
 # ---- Brand Memory, shown during setup ----
 P('ao-24', 'onboarding', 'about', 'orchestrator', editable='yes', title='About the brand',
   description="Angel One calls itself India's leading stock broker and a one-stop investment platform: 3.8 Cr+ registered users, 10 Cr+ app downloads and 30+ years of trust. Stocks, F&O, mutual funds, IPOs, commodities and US stocks sit in one app.",
   tags=['Digital-first', 'Retail investors', 'One app']),
 P('ao-25', 'onboarding', 'brandkit', 'orchestrator'),
 P('ao-26', 'onboarding', 'about', 'orchestrator', editable='yes', title='Voice and tone',
   description="Plain, number-led and reassuring. Short benefit lines like 'Transparent pricing. No hidden charges' and 'Where India invests', the rupee figure up front, and the SEBI disclaimer under every claim.",
   tags=['Number-led', 'Plain', 'Compliant']),
 P('ao-27', 'onboarding', 'connect', 'orchestrator', connect='wordpress'),
 P('ao-28', 'onboarding', 'about', 'orchestrator', editable='yes', title='Company',
   description='Angel One Limited, formerly Angel Broking, founded in 1987 and registered in Andheri East, Mumbai. A SEBI-registered broker, research analyst and investment adviser, a CDSL depository participant and an AMFI mutual fund distributor.',
   tags=['Mumbai', 'Since 1987', 'Listed', 'NSE, BSE, MCX, NCDEX', '10,000+ authorised persons', '+Add team size']),
 P('ao-29', 'onboarding', 'about', 'orchestrator', editable='yes', title='Product catalog',
   description='Trading and investing products sold in one app, with 5,877 stock pages, 7,266 mutual fund pages, 1,292 IPO pages, 1,471 US stock pages and 243 ETF pages on the site.',
   tags=['Stocks', 'F&O', 'Mutual Funds', 'SIP', 'IPO', 'Commodities', 'US Stocks', 'ETF', 'MTF', 'Fixed Deposits']),
 P('ao-30', 'onboarding', 'about', 'orchestrator', editable='yes', title='Product information',
   description='Offers on the site today: ₹0 account opening, ₹0 AMC for the first year then ₹20 a month, and ₹0 brokerage for 30 days up to ₹500. After that, stocks cost the lower of ₹20 or 0.1% an order and F&O ₹20. SIPs start at ₹100 with zero commission.'),
 P('ao-31', 'onboarding', 'about', 'orchestrator', editable='yes', title='Audience',
   description='First-time investors opening a demat account on their phone, SIP savers starting at ₹100, and active traders who want fast orders, advanced charts, TradingView and SmartAPI. 10,000+ authorised persons sell it well beyond the metros.'),
 P('ao-32', 'onboarding', 'about', 'orchestrator', editable='yes', title='Industry leaders',
   tags=['zerodha.com #387ed1', 'groww.in #00d09c', 'upstox.com #5a298b', 'dhan.co #1a1a1a']),
]

def main():
    wb = Workbook(); ws = wb.active; ws.title = 'AngelOne-1'
    bold = Font(bold=True); head = PatternFill('solid', fgColor='E8ECFF')
    for label, key, hint in BRAND_KEYS:
        ws.append([label, BRAND.get(key, ''), hint])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append([c for c, _ in COLUMNS])
    for c in ws[ws.max_row]: c.font = bold; c.fill = head
    for p in POSTS:
        ws.append([p.get(c, '') for c, _ in COLUMNS])
    for col, w in zip('ABCDEFGHIJKLMNOPQRSTUVWXYZ', [12, 10, 16, 10, 13, 12, 8, 44, 70, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    out = os.path.join(ROOT, 'data', 'angelone-seed.xlsx')
    wb.save(out); print('Wrote', os.path.relpath(out, ROOT), len(POSTS), 'posts')

if __name__ == '__main__': main()
